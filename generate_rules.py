from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_participant_rules():
    doc = Document()
    
    # Title
    title = doc.add_heading('CYBER VERSE: MISSION RULES', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Intro
    p = doc.add_paragraph()
    run = p.add_run('Welcome, Agents. Below are the official operational protocols for the Cyber Verse decryption challenge.')
    run.font.size = Pt(12)
    
    # Section 1: Objective
    doc.add_heading('1. Objective', level=1)
    doc.add_paragraph('Your team must decrypt the central network grid by correctly identifying 24 cybersecurity-related terms based on the encrypted hints provided at each node.')

    # Section 2: Scoring & Penalties
    doc.add_heading('2. Scoring & Penalties', level=1)
    
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Action'
    hdr_cells[1].text = 'Point Impact'
    
    row_cells = table.add_row().cells
    row_cells[0].text = 'Correct Decryption'
    row_cells[1].text = '+100 Points'
    
    row_cells = table.add_row().cells
    row_cells[0].text = 'Incorrect Attempt'
    row_cells[1].text = '-2 Points'
    
    row_cells = table.add_row().cells
    row_cells[0].text = 'Activating a Hint'
    row_cells[1].text = '-60 Points'

    # Section 3: The Hint Mechanism
    doc.add_heading('3. The Hint Mechanism', level=1)
    p = doc.add_paragraph()
    p.add_run('Stuck on a node? ').bold = True
    p.add_run('After 10 failed attempts, the "Hint" protocol becomes available. Activating it will reveal a Python script. You must analyze the logic of this script to determine the secret key (answer). Use hints sparingly, as each one carries a heavy point penalty.')

    # Section 4: Rules of Engagement
    doc.add_heading('4. Rules of Engagement', level=1)
    doc.add_paragraph('• Do not refresh your browser once the mission has started.', style='List Bullet')
    doc.add_paragraph('• Avoid switching tabs or minimizing the game window; the system tracks active focus.', style='List Bullet')
    doc.add_paragraph('• Collaboration within your assigned team is encouraged, but cross-team communication is strictly prohibited.', style='List Bullet')
    doc.add_paragraph('• The mission ends when the entire grid is secured. Your final rank is determined by Score, then by Completion Time.', style='List Bullet')

    # Footer
    doc.add_paragraph('\n---')
    footer = doc.add_paragraph('STATUS: CLASSIFIED | PREPARE FOR DEPLOYMENT')
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save('Participant_Rules.docx')
    print("Successfully generated Participant_Rules.docx")

if __name__ == '__main__':
    create_participant_rules()
